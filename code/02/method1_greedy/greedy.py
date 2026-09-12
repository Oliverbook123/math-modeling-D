# -*- coding: utf-8 -*-
"""
问题 2 方案 1：贪心 + 优先级（冲突图贪心着色）

运行方式（项目根目录下，强制用 uv，禁止 pip / 裸 python）：
    UV_INDEX_URL="https://mirrors.aliyun.com/pypi/simple/" uv run --with openpyxl python code/02/method1_greedy/main.py

算法概述（docs/plan/方案1_贪心优先级.md）：
  1. 用 code/01/common/data_loader.py 的 load_plans() 读附件1（150 个计划）；
  2. 重算全部冲突对（应与问题 1 的 237 对一致，不一致会打印说明）；
  3. 排序键：(类别 A<B<C, 冲突度数降序, 编号)；
  4. 依序"入座"：每个计划先试保留（与全部已保留计划 0 冲突即保留），
     否则按 幅度|Δ| 从 1 递增 枚举动作，每档内顺序：频段-Δ、频段+Δ、时间-Δ、时间+Δ
     （频段 |Δ|<=10，时间 |Δ|<=5，平移后频段仍须 ⊆[0,100)，时间起点 >=0）；
     每个候选动作都与 全部已保留（含已调整）计划 重新判冲突；
     全部候选都不行 → 撤销；
  5. 输出 result.csv / summary.json / result2_greedy.xlsx 到 code/output/02/method1_greedy/；
  6. 尾部执行 README §6 全部自检，assert 失败即非 0 退出。
"""

import csv
import json
import os
import sys
from itertools import combinations

# ---------- sys.path 注入，复用问题 1 的公共数据模块（不得修改 code/01/） ----------
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))          # .../code/02/method1_greedy
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(CURRENT_DIR)))  # 项目根
CODE01_DIR = os.path.join(PROJECT_ROOT, "code", "01")
if CODE01_DIR not in sys.path:
    sys.path.insert(0, CODE01_DIR)

from common.data_loader import Plan, load_plans, get_data_path  # noqa: E402

import openpyxl  # 由 uv --with openpyxl 提供

# ---------- 常量 ----------
MAX_FREQ_SHIFT = 10   # 频段最大平移 |Δf|
MAX_TIME_SHIFT = 5    # 时间最大平移 |Δt|
FREQ_BAND = (0, 100)  # 频段可用范围 [0,100)
EXPECTED_CONFLICTS = 237  # 问题 1 检出的冲突对数（基准）

CATEGORY_RANK = {"A": 0, "B": 1, "C": 2}  # 优先级排序用：A 最先处理

# 输出目录：code/output/02/method1_greedy/
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "code", "output", "02", "method1_greedy")


# ======================================================================
# 冲突判定（与问题 1 完全同口径：频段交叠 ∧ 任一时间实例交叠）
# 区间 [a,b) 与 [c,d) 交叠 ⟺ a < d 且 c < b
# ======================================================================
def time_overlap(p1: Plan, p2: Plan) -> bool:
    """p1 的任一时间实例与 p2 的任一时间实例交叠则返回 True"""
    for s1, e1 in p1.time_instances():
        for s2, e2 in p2.time_instances():
            if s1 < e2 and s2 < e1:  # 左闭右开区间交叠条件
                return True
    return False


def conflict(a: Plan, b: Plan) -> bool:
    """两个计划是否冲突：频段交叠 且 存在交叠的时间实例"""
    if not (a.freq_start < b.freq_end and b.freq_start < a.freq_end):
        return False  # 频段不相交，快速跳过
    return time_overlap(a, b)


def find_all_conflicts(plans: list[Plan]) -> list[tuple[str, str]]:
    """全量两两冲突检测，返回冲突对 (id1, id2) 列表（id1 在输入序中在前）"""
    pairs = []
    for a, b in combinations(plans, 2):
        if conflict(a, b):
            pairs.append((a.id, b.id))
    return pairs


# ======================================================================
# 候选动作枚举与生成
# ======================================================================
def shifted_freq(p: Plan, d: int) -> Plan:
    """频段整体平移 d 后的新 Plan（时间/次数/间隔不变）"""
    return Plan(p.id, p.category, p.freq_start + d, p.freq_end + d,
                p.time_start, p.time_end, p.interval, p.count)


def shifted_time(p: Plan, d: int) -> Plan:
    """首次时间区间整体平移 d（所有实例随动）后的新 Plan"""
    return Plan(p.id, p.category, p.freq_start, p.freq_end,
                p.time_start + d, p.time_end + d, p.interval, p.count)


def enumerate_actions(p: Plan):
    """
    生成候选动作生成器，顺序（方案文档规定）：
      |Δ| 从 1 递增；每档内：频段-Δ、频段+Δ、时间-Δ、时间+Δ
    每个动作 = (决策, Δ频段, Δ时间, 平移后的 Plan)，且已过滤越界动作：
      平移后频段须仍落在 [0,100)，时间起点须 >= 0
    """
    for delta in range(1, max(MAX_FREQ_SHIFT, MAX_TIME_SHIFT) + 1):
        # 频段 -delta / +delta（|Δf| <= 10，平移后 ⊆ [0,100)）
        if delta <= MAX_FREQ_SHIFT:
            for d in (-delta, delta):
                if p.freq_start + d >= FREQ_BAND[0] and p.freq_end + d <= FREQ_BAND[1]:
                    yield ("调频段", d, 0, shifted_freq(p, d))
        # 时间 -delta / +delta（|Δt| <= 5，平移后首次起点 >= 0）
        if delta <= MAX_TIME_SHIFT:
            for d in (-delta, delta):
                if p.time_start + d >= 0:
                    yield ("调时间", 0, d, shifted_time(p, d))


# ======================================================================
# 贪心主流程
# ======================================================================
def solve(plans: list[Plan]):
    """
    按 (类别 A<B<C, 冲突度数降序, 编号) 依序处理。
    返回 decisions: dict[id] = {"决策", "Δ频段", "Δ时间", "plan": 最终 Plan 或 None}
    """
    # ---- 冲突度数（无向图度数），用于同类内排序 ----
    degree = {p.id: 0 for p in plans}
    for a, b in combinations(plans, 2):
        if conflict(a, b):
            degree[a.id] += 1
            degree[b.id] += 1

    # ---- 排序：优先级 A>B>C 即处理顺序，不引入任何打乱该顺序的启发式 ----
    order = sorted(plans, key=lambda p: (CATEGORY_RANK[p.category], -degree[p.id], p.id))

    kept: list[Plan] = []   # 已保留（含已调整）的计划集合
    decisions = {}
    for p in order:
        # 1) 先试"保留"：与全部已保留计划无冲突则原样保留
        if not any(conflict(p, q) for q in kept):
            decisions[p.id] = {"决策": "保留", "Δ频段": 0, "Δ时间": 0, "plan": p}
            kept.append(p)
            continue
        # 2) 有冲突：按 |Δ| 从小到大枚举平移动作，
        #    每个候选动作都与【全部已保留（含已调整）计划】重新判冲突
        done = False
        for name, df, dt, newp in enumerate_actions(p):
            if not any(conflict(newp, q) for q in kept):
                decisions[p.id] = {"决策": name, "Δ频段": df, "Δ时间": dt, "plan": newp}
                kept.append(newp)
                done = True
                break
        # 3) 所有平移都不行 → 撤销
        if not done:
            decisions[p.id] = {"决策": "撤销", "Δ频段": 0, "Δ时间": 0, "plan": None}
    return decisions, order


# ======================================================================
# 输出三件套
# ======================================================================
def fmt_interval(a: int, b: int) -> str:
    """区间格式，同附件1：左闭右开 [a,b)"""
    return f"[{a},{b})"


def write_result_csv(plans, decisions):
    """result.csv：全部 150 个计划的决策明细（UTF-8 with BOM）"""
    path = os.path.join(OUTPUT_DIR, "result.csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["装备编号", "类别", "决策(保留/调频段/调时间/撤销)",
                    "频段平移", "时间平移", "调整后频段区间", "调整后时间区间"])
        for p in plans:  # 按附件1原顺序输出
            d = decisions[p.id]
            if d["决策"] == "撤销":
                fq, tq = "", ""  # 撤销不写区间
            else:
                q = d["plan"]
                fq, tq = fmt_interval(q.freq_start, q.freq_end), fmt_interval(q.time_start, q.time_end)
            w.writerow([p.id, p.category, d["决策"], d["Δ频段"], d["Δ时间"], fq, tq])
    return path


def write_result2_xlsx(plans, decisions):
    """result2_greedy.xlsx：严格按附件2/result2.xlsx 模板，只写被动过的装备"""
    path = os.path.join(OUTPUT_DIR, "result2_greedy.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # 表头 4 列逐字对齐模板
    ws.append(["用频装备编号", "调整后频段区间", "调整后时间区间", "是否撤销用频计划"])
    for p in plans:
        d = decisions[p.id]
        if d["决策"] == "保留":
            continue  # 未动的装备一行都不写
        if d["决策"] == "撤销":
            ws.append([p.id, None, None, "是"])
        elif d["决策"] == "调频段":
            q = d["plan"]
            ws.append([p.id, fmt_interval(q.freq_start, q.freq_end), None, None])
        elif d["决策"] == "调时间":
            q = d["plan"]
            # 调时间 → C 列写平移后的【首次】时间区间（实例随动）
            ws.append([p.id, None, fmt_interval(q.time_start, q.time_end), None])
    wb.save(path)
    return path


# ======================================================================
# 主程序
# ======================================================================
def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # ---------- 1. 读数据 ----------
    plans = load_plans(get_data_path())
    assert len(plans) == 150, f"计划数应为150，实际 {len(plans)}"
    cat_count = {}
    for p in plans:
        cat_count[p.category] = cat_count.get(p.category, 0) + 1
    assert cat_count == {"A": 20, "B": 40, "C": 90}, f"类别数量异常: {cat_count}"

    # ---------- 2. 重算冲突，对基准 237 ----------
    base_conflicts = find_all_conflicts(plans)
    n_base = len(base_conflicts)
    print(f"[基准] 重算冲突对数 = {n_base}（问题1基准 = {EXPECTED_CONFLICTS}）")
    if n_base != EXPECTED_CONFLICTS:
        print(f"[警告] 与问题1的 237 对不一致，差值 {n_base - EXPECTED_CONFLICTS}，请核对冲突口径")

    # A-A 冲突对（问题1 中恰 1 对），单独跟踪其处理方式
    aa_pairs = [(a, b) for a, b in base_conflicts if a[0] == "A" and b[0] == "A"]
    print(f"[基准] A-A 冲突对数 = {len(aa_pairs)}: {aa_pairs}")

    # ---------- 3. 贪心求解 ----------
    decisions, order = solve(plans)
    # 记录每个冲突对中实际先被处理的一方（用于 A-A 对的说明）
    pos = {p.id: i for i, p in enumerate(order)}
    first_processed = {f"{a}|{b}": (a if pos[a] < pos[b] else b) for a, b in aa_pairs}

    # ---------- 4. 表1 统计 ----------
    table1 = {c: {"保留": 0, "调整": 0, "撤销": 0} for c in "ABC"}
    for p in plans:
        d = decisions[p.id]["决策"]
        key = "调整" if d in ("调频段", "调时间") else d
        table1[p.category][key] += 1

    n_revoke = sum(1 for p in plans if decisions[p.id]["决策"] == "撤销")
    n_adjust = sum(1 for p in plans if decisions[p.id]["决策"] in ("调频段", "调时间"))
    total_shift = sum(abs(decisions[p.id]["Δ频段"]) + abs(decisions[p.id]["Δ时间"]) for p in plans)

    # A-A 冲突对的处理记录（方案文档特别要求）
    aa_handling = []
    for a, b in aa_pairs:
        first = first_processed[a + "|" + b]
        second = b if first == a else a
        aa_handling.append({
            "冲突对": [a, b],
            "先处理者": first,  # 排序键：同类内冲突度数降序、再编号升序
            a: decisions[a]["决策"],
            b: decisions[b]["决策"],
            "先处理者决策": decisions[first]["决策"],
            "后处理者决策": decisions[second]["决策"],
            "说明": "排序后先处理者保留，后处理者平移消解" if decisions[second]["决策"] != "撤销"
                    else "后处理者被迫撤销（应在 summary 中解释）",
        })

    # ---------- 5. 写产出文件 ----------
    csv_path = write_result_csv(plans, decisions)
    xlsx_path = write_result2_xlsx(plans, decisions)

    # ---------- 6. README §6 自检（全部 assert，失败即非 0 退出） ----------
    checks = {}

    # 6.1 最终方案（保留+调整，剔除撤销）重新冲突检测 == 0
    final_plans = [d["plan"] for d in decisions.values() if d["plan"] is not None]
    final_conflicts = find_all_conflicts(final_plans)
    assert len(final_conflicts) == 0, f"自检失败：最终仍有冲突 {len(final_conflicts)} 对: {final_conflicts[:5]}"
    checks["最终冲突数"] = 0

    # 6.2/6.3 每个被动计划：只动 1 个参数、幅度不超限、频段⊆[0,100)、起点>=0、次数/间隔不变
    orig = {p.id: p for p in plans}
    n_params_changed_max = 0
    for p in plans:
        d = decisions[p.id]
        if d["决策"] == "撤销":
            continue
        q = d["plan"]
        df, dt = d["Δ频段"], d["Δ时间"]
        assert not (df != 0 and dt != 0), f"{p.id}: 同时动了两个参数"
        n_params_changed_max = max(n_params_changed_max, 1 if (df or dt) else 0)
        assert abs(df) <= MAX_FREQ_SHIFT, f"{p.id}: |Δf|={abs(df)} 超限"
        assert abs(dt) <= MAX_TIME_SHIFT, f"{p.id}: |Δt|={abs(dt)} 超限"
        assert q.freq_start >= FREQ_BAND[0] and q.freq_end <= FREQ_BAND[1], f"{p.id}: 频段越界"
        assert q.time_start >= 0, f"{p.id}: 时间起点为负"
        # 平移一致性：调频段只动频段两端、调时间只动时间两端，宽度/次数/间隔不变
        if d["决策"] == "调频段":
            assert (q.freq_start, q.freq_end) == (p.freq_start + df, p.freq_end + df)
            assert (q.time_start, q.time_end) == (p.time_start, p.time_end)
        elif d["决策"] == "调时间":
            assert (q.time_start, q.time_end) == (p.time_start + dt, p.time_end + dt)
            assert (q.freq_start, q.freq_end) == (p.freq_start, p.freq_end)
        else:  # 保留：原样不动
            assert (q.freq_start, q.freq_end, q.time_start, q.time_end) == \
                   (p.freq_start, p.freq_end, p.time_start, p.time_end)
        assert q.interval == p.interval and q.count == p.count, f"{p.id}: 间隔/次数被改动"
    assert n_params_changed_max <= 1
    checks["每计划最多动1参数"] = True

    # 6.4 表1 三行总数 = 20/40/90
    assert sum(table1["A"].values()) == 20
    assert sum(table1["B"].values()) == 40
    assert sum(table1["C"].values()) == 90
    checks["表1总数20_40_90"] = True

    # 6.5 result2_greedy.xlsx 读回验证：行数 = 调整数 + 撤销数，格式列对齐
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb.active
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    assert len(rows) == n_adjust + n_revoke, \
        f"xlsx 行数 {len(rows)} != 调整数+撤销数 {n_adjust + n_revoke}"
    touched = [p for p in plans if decisions[p.id]["决策"] != "保留"]
    assert len(rows) == len(touched)
    seen_ids = set()
    for row, p in zip(rows, touched):
        eid, fq, tq, rv = row
        assert eid == p.id, f"xlsx 行 {eid} 与决策序 {p.id} 不一致"
        seen_ids.add(eid)
        d = decisions[p.id]
        if d["决策"] == "撤销":
            assert fq is None and tq is None and rv == "是", f"{eid}: 撤销行列不对"
        elif d["决策"] == "调频段":
            q = d["plan"]
            assert fq == fmt_interval(q.freq_start, q.freq_end) and tq is None and rv is None, \
                f"{eid}: 调频段行列不对"
            assert fq.startswith("[") and fq.endswith(")"), f"{eid}: 区间格式非 [a,b)"
        else:
            q = d["plan"]
            assert tq == fmt_interval(q.time_start, q.time_end) and fq is None and rv is None, \
                f"{eid}: 调时间行列不对"
            assert tq.startswith("[") and tq.endswith(")"), f"{eid}: 区间格式非 [a,b)"
    assert seen_ids == {p.id for p in touched}
    checks["xlsx读回行数=调整+撤销且列对齐"] = True

    # ---------- 7. summary.json ----------
    summary = {
        "方法": "方案1_贪心优先级（冲突图贪心着色）",
        "表1": table1,
        "目标函数": {
            "第1层_撤销总数": n_revoke,
            "第2层_A撤销": table1["A"]["撤销"],
            "第2层_B撤销": table1["B"]["撤销"],
            "第2层_C撤销": table1["C"]["撤销"],
            "第3层_调整数": n_adjust,
            "第4层_总调整幅度": total_shift,
        },
        "冲突对数_原始": n_base,
        "冲突对数_原始_与问题1基准一致": n_base == EXPECTED_CONFLICTS,
        "A_A冲突对": aa_handling,
        "自检": checks,
    }
    sp = os.path.join(OUTPUT_DIR, "summary.json")
    with open(sp, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print("全部自检通过 ✔")
    print("产出:", csv_path, sp, xlsx_path, sep="\n  ")


if __name__ == "__main__":
    main()
