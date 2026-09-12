# -*- coding: utf-8 -*-
"""
问题3 公共模块：数据加载与工具函数
负责：加载问题2结果、构建占用集合、C类装备放置检测等
"""

import os
import re
import sys
import importlib.util
from typing import Set, Tuple, List

# ---------- 从 code/01/common/data_loader.py 动态加载 Plan 等 ----------
# 不能直接 from common.data_loader import ... 因为 code/03/common/ 会和 code/01/common/ 冲突
_current_dir = os.path.dirname(os.path.abspath(__file__))
_project_root = os.path.dirname(os.path.dirname(os.path.dirname(_current_dir)))
_code01_loader = os.path.join(_project_root, "code", "01", "common", "data_loader.py")

_spec = importlib.util.spec_from_file_location("_code01_data_loader", _code01_loader)
_mod = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_mod)

Plan = _mod.Plan
load_plans = _mod.load_plans
get_data_path = _mod.get_data_path

import openpyxl

# ---------- 常量 ----------
FREQ_BAND = (0, 100)
C_FREQ_WIDTH = 3
C_TIME_WIDTH = 2
C_INTERVAL = 8
C_COUNT = 12


def parse_interval(s: str) -> Tuple[int, int]:
    """解析区间字符串 '[a,b)' -> (a, b)"""
    s = s.strip()
    match = re.match(r'[\[\(]\s*(\d+)\s*,\s*(\d+)\s*[\]\)]', s)
    if not match:
        raise ValueError(f"无法解析区间: {s}")
    return int(match.group(1)), int(match.group(2))


def fmt_interval(a: int, b: int) -> str:
    """格式化区间 [a,b)"""
    return f"[{a},{b})"


def get_project_root() -> str:
    """获取项目根目录"""
    return _project_root


def load_problem2_result(result2_path: str, plans: List[Plan]) -> List[Plan]:
    """
    从 result2_xxx.xlsx 读取问题2的消解结果，返回最终无冲突的活跃计划。
    撤销(D列="是")的跳过，有调整的用调整后区间，未出现的原样保留。
    """
    orig = {p.id: p for p in plans}
    wb = openpyxl.load_workbook(result2_path, read_only=True)
    ws = wb.active

    adjusted_ids, revoked_ids = set(), set()
    adjusted_plans = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        pid = str(row[0]).strip()
        freq_str, time_str, revoked = row[1], row[2], row[3]

        if revoked and str(revoked).strip() == "是":
            revoked_ids.add(pid)
            continue

        adjusted_ids.add(pid)
        p = orig[pid]
        fs, fe = parse_interval(str(freq_str)) if freq_str and str(freq_str).strip() else (p.freq_start, p.freq_end)
        ts, te = parse_interval(str(time_str)) if time_str and str(time_str).strip() else (p.time_start, p.time_end)

        adjusted_plans[pid] = Plan(
            id=p.id, category=p.category,
            freq_start=fs, freq_end=fe,
            time_start=ts, time_end=te,
            interval=p.interval, count=p.count
        )
    wb.close()

    active = []
    for p in plans:
        if p.id in revoked_ids:
            continue
        active.append(adjusted_plans.get(p.id, p))
    return active


def build_occupied_set(plans: List[Plan]) -> Set[Tuple[int, int]]:
    """构建已占用的 (频段, 时间) 格子集合"""
    occupied = set()
    for p in plans:
        for ts, te in p.time_instances():
            for f in range(p.freq_start, p.freq_end):
                for t in range(ts, te):
                    occupied.add((f, t))
    return occupied


def get_max_time(plans: List[Plan]) -> int:
    """获取所有计划中最大时间值"""
    mx = 0
    for p in plans:
        for ts, te in p.time_instances():
            mx = max(mx, te)
    return mx


def make_c_plan(freq_start: int, time_start: int) -> Plan:
    """在指定位置创建一个 C 类装备计划"""
    return Plan(
        id="C_NEW", category="C",
        freq_start=freq_start, freq_end=freq_start + C_FREQ_WIDTH,
        time_start=time_start, time_end=time_start + C_TIME_WIDTH,
        interval=C_INTERVAL, count=C_COUNT
    )


def can_place_c(freq_start: int, time_start: int,
                occupied: Set[Tuple[int, int]]) -> bool:
    """检查 C 类装备能否放在 (freq_start, time_start)，所有12次使用都不冲突"""
    if freq_start < 0 or freq_start + C_FREQ_WIDTH > FREQ_BAND[1]:
        return False
    if time_start < 0:
        return False
    for k in range(C_COUNT):
        ts = time_start + k * C_INTERVAL
        te = ts + C_TIME_WIDTH
        for f in range(freq_start, freq_start + C_FREQ_WIDTH):
            for t in range(ts, te):
                if (f, t) in occupied:
                    return False
    return True


def place_c(freq_start: int, time_start: int,
            occupied: Set[Tuple[int, int]]) -> Plan:
    """放置 C 类装备并更新占用集合"""
    p = make_c_plan(freq_start, time_start)
    for ts, te in p.time_instances():
        for f in range(p.freq_start, p.freq_end):
            for t in range(ts, te):
                occupied.add((f, t))
    return p


def plans_conflict(p1: Plan, p2: Plan) -> bool:
    """两个计划是否冲突：频段交叠 ∧ 时间实例交叠"""
    if not (p1.freq_start < p2.freq_end and p2.freq_start < p1.freq_end):
        return False
    for ts1, te1 in p1.time_instances():
        for ts2, te2 in p2.time_instances():
            if ts1 < te2 and ts2 < te1:
                return True
    return False


def verify_no_conflict(plans: List[Plan]) -> int:
    """验证计划列表无冲突，返回冲突对数"""
    n, cnt = len(plans), 0
    for i in range(n):
        for j in range(i + 1, n):
            if plans_conflict(plans[i], plans[j]):
                cnt += 1
    return cnt


def write_result3_xlsx(new_plans: List[Plan], output_path: str):
    """
    写入 result3.xlsx（严格按模板格式）
    A列=新增用频装备序号, B列=调整后频段区间, C列=调整后时间区间
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["新增用频装备序号", "调整后频段区间", "调整后时间区间"])
    for i, p in enumerate(new_plans, 1):
        ws.append([i, fmt_interval(p.freq_start, p.freq_end),
                   fmt_interval(p.time_start, p.time_end)])
    wb.save(output_path)
